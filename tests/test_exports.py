"""Testes do core.exports — DB de fixture em tmp_path, sem rede, sem segredo."""
import csv
import json
import sqlite3

import pytest

from core import database
from core.exports import export_table, list_kinds, list_tables


@pytest.fixture()
def db(tmp_path, monkeypatch):
    """Banco de fixture com panel_entities/products/raw_snapshots povoados."""
    path = str(tmp_path / "fixture.db")
    monkeypatch.setattr(database, "DB_PATH", path)
    database.init_db()
    database.save_entities("blackbr.customer", [
        ("1", {"username": "alice", "expires_at": "2027-01-01T02:59:59Z"}),
        ("2", {"username": "bob", "expires_at": "2027-02-01T02:59:59Z"}),
    ])
    database.save_entities("blackbr.package", [("9", {"name": "Pack", "connections": 2})])
    database.save_raw_and_clean("https://x/item/1", "1", "Item 1", 9.9, {"a": 1})
    return path


def test_list_tables_e_kinds(db):
    assert {"raw_snapshots", "products", "panel_entities"} <= set(list_tables(db))
    assert set(list_kinds(db)) == {"blackbr.customer", "blackbr.package"}


def test_export_panel_entities_csv(db, tmp_path):
    res = export_table(out_path=str(tmp_path / "e.csv"), db_path=db)
    assert res["rows"] == 3
    assert res["format"] == "csv"
    with open(res["path"], newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    assert {r["kind"] for r in rows} == {"blackbr.customer", "blackbr.package"}
    alice = next(r for r in rows if r["external_id"] == "1")
    assert alice["username"] == "alice"  # payload achatado em coluna


def test_export_filtro_kind_e_site(db, tmp_path):
    res_kind = export_table(kind="blackbr.customer", db_path=db,
                            out_path=str(tmp_path / "k.csv"))
    res_site = export_table(site="blackbr", db_path=db,
                            out_path=str(tmp_path / "s.csv"))
    assert res_kind["rows"] == 2
    assert res_site["rows"] == 3
    assert res_site["path"].endswith(".csv")


def test_export_fields_subconjunto_e_erro(db, tmp_path):
    res = export_table(fields=["kind", "username"], db_path=db,
                       out_path=str(tmp_path / "f.csv"))
    assert res["columns"] == ["kind", "username"]
    with pytest.raises(ValueError, match="inexistentes"):
        export_table(fields=["password_do_nada"], db_path=db,
                     out_path=str(tmp_path / "x.csv"))


def test_export_json(db, tmp_path):
    res = export_table(fmt="json", kind="blackbr.customer", db_path=db,
                       out_path=str(tmp_path / "e.json"))
    data = json.loads(open(res["path"], encoding="utf-8").read())
    assert len(data) == 2
    assert data[0]["kind"] == "blackbr.customer"


def test_export_xlsx(db, tmp_path):
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook

    res = export_table(fmt="xlsx", db_path=db, out_path=str(tmp_path / "e.xlsx"))
    ws = load_workbook(res["path"]).active
    linhas = list(ws.iter_rows(values_only=True))
    assert linhas[0][0] == "kind"
    assert len(linhas) == 4  # header + 3


def test_export_products_e_raw(db, tmp_path):
    res_p = export_table(table="products", db_path=db,
                         out_path=str(tmp_path / "p.json"))
    assert json.loads(open(res_p["path"], encoding="utf-8").read())[0]["title"] == "Item 1"
    res_r = export_table(table="raw_snapshots", db_path=db,
                         out_path=str(tmp_path / "r.csv"))
    assert res_r["rows"] == 1


def test_export_tabela_desconhecida_e_injecao(db, tmp_path):
    with pytest.raises(ValueError, match="não existe"):
        export_table(table="nope", db_path=db)
    with pytest.raises(ValueError, match="inválido"):
        export_table(table="panel_entities; DROP--", db_path=db)


def test_export_limit_e_out_default(db, tmp_path):
    res = export_table(limit=1, db_path=db, out_dir=tmp_path)
    assert res["rows"] == 1
    assert str(tmp_path) in res["path"]  # default nameado em out_dir
    assert "export_panel_entities_all_" in res["path"]


def test_export_kind_outra_tabela_erro(db, tmp_path):
    with pytest.raises(ValueError, match="só se aplicam"):
        export_table(table="products", kind="x", db_path=db)
